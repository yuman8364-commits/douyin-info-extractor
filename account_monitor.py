# -*- coding: utf-8 -*-
"""抖音账号公开主页监控（不读取系统剪贴板）。"""

from __future__ import annotations

from dataclasses import dataclass
import csv
import html
import json
import re
from typing import Any
from urllib.parse import quote, unquote

import requests
from openpyxl import load_workbook


USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)
PROFILE_RE = re.compile(r"douyin\.com/user/([^/?#\s]+)", re.I)
ACCOUNT_COLUMN_NAMES = {
    "抖音号", "抖音id", "抖音账号", "账号id", "账号", "主页链接",
    "抖音主页", "抖音主页链接", "secuid", "sec_uid",
}


class MonitorError(RuntimeError):
    """账号检查无法得出可信结论。"""


class AccountAbnormalError(MonitorError):
    """公开页面明确表明账号不存在或状态异常。"""


@dataclass(frozen=True)
class AccountSnapshot:
    account_id: str
    nickname: str
    latest_aweme_id: str
    latest_create_time: int
    latest_desc: str


@dataclass(frozen=True)
class AccountImportResult:
    accounts: tuple[str, ...]
    invalid_cells: tuple[str, ...]
    matched_columns: tuple[str, ...]


def _normalized_header(value) -> str:
    return re.sub(r"[\s_\-]+", "", str(value or "")).casefold()


def _import_values(rows, source_name: str) -> AccountImportResult:
    materialized = [tuple(row) for row in rows]
    aliases = {_normalized_header(name) for name in ACCOUNT_COLUMN_NAMES}
    header_index = None
    columns: list[int] = []
    headers: list[str] = []
    for row_index, row in enumerate(materialized[:20]):
        matched = [index for index, value in enumerate(row) if _normalized_header(value) in aliases]
        if matched:
            header_index = row_index
            columns = matched
            headers = [f"{source_name}!{row[index]}" for index in matched]
            break
    if header_index is None:
        return AccountImportResult((), (), ())

    accounts: list[str] = []
    invalid: list[str] = []
    for row_number, row in enumerate(materialized[header_index + 1 :], start=header_index + 2):
        for column in columns:
            value = row[column] if column < len(row) else None
            if value is None or not str(value).strip():
                continue
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            try:
                account_id = normalize_account_input(str(value))
            except ValueError:
                invalid.append(f"{source_name}!R{row_number}C{column + 1}")
                continue
            if account_id not in accounts:
                accounts.append(account_id)
    return AccountImportResult(tuple(accounts), tuple(invalid), tuple(headers))


def import_accounts_from_table(path) -> AccountImportResult:
    """只读 Excel/CSV/TSV，从常见账号列批量导入并稳定去重。"""
    suffix = str(path).lower().rsplit(".", 1)[-1]
    results: list[AccountImportResult] = []
    if suffix in {"xlsx", "xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                results.append(_import_values(sheet.iter_rows(values_only=True), sheet.title))
        finally:
            workbook.close()
    elif suffix in {"csv", "tsv"}:
        delimiter = "\t" if suffix == "tsv" else ","
        with open(path, "r", encoding="utf-8-sig", newline="") as stream:
            results.append(_import_values(csv.reader(stream, delimiter=delimiter), str(path)))
    else:
        raise ValueError("只支持 .xlsx、.xlsm、.csv 或 .tsv 表格")

    accounts: list[str] = []
    invalid: list[str] = []
    matched: list[str] = []
    for result in results:
        for account_id in result.accounts:
            if account_id not in accounts:
                accounts.append(account_id)
        invalid.extend(result.invalid_cells)
        matched.extend(result.matched_columns)
    return AccountImportResult(tuple(accounts), tuple(invalid), tuple(matched))


def is_new_work(previous: dict, current: AccountSnapshot) -> bool:
    """只有已有基线且最新作品确实前进时才判定为新发布。"""
    if not previous.get("baseline_initialized") or not current.latest_aweme_id:
        return False
    previous_id = str(previous.get("latest_aweme_id") or "")
    previous_time = int(previous.get("latest_create_time") or 0)
    return current.latest_create_time > previous_time or (
        current.latest_create_time == previous_time
        and current.latest_aweme_id != previous_id
    )


def normalize_account_input(value: str) -> str:
    """接受抖音号、sec_uid 或主页链接，返回用于查询的账号标识。"""
    text = (value or "").strip()
    match = PROFILE_RE.search(text)
    if match:
        text = unquote(match.group(1))
    text = text.strip().lstrip("@")
    if not text or len(text) > 160 or any(char.isspace() for char in text):
        raise ValueError("请输入抖音号、sec_uid 或完整主页链接")
    return text


def _json_candidates(page: str):
    for pattern in (
        r'<script[^>]+id=["\']RENDER_DATA["\'][^>]*>(.*?)</script>',
        r'<script[^>]+id=["\']__ROUTER_DATA__["\'][^>]*>(.*?)</script>',
        r'<script[^>]+type=["\']application/json["\'][^>]*>(.*?)</script>',
    ):
        for match in re.finditer(pattern, page, re.I | re.S):
            raw = html.unescape(match.group(1).strip())
            for candidate in (raw, unquote(raw)):
                try:
                    yield json.loads(candidate)
                    break
                except (TypeError, json.JSONDecodeError):
                    continue


def _walk(value: Any):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from _walk(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk(child)


def _account_matches(user: dict, account_id: str) -> bool:
    expected = account_id.casefold()
    values = (
        user.get("unique_id"), user.get("short_id"), user.get("uid"),
        user.get("sec_uid"), user.get("secUid"),
    )
    return any(str(value or "").casefold() == expected for value in values)


def parse_account_snapshot(page: str, account_id: str) -> AccountSnapshot:
    """从搜索/主页 SSR 数据中提取账号及最新公开作品。"""
    lowered = page.lower()
    if any(marker in page for marker in ("帐号已注销", "账号已注销", "用户不存在")):
        raise AccountAbnormalError("账号已注销或不存在")
    if "ttgcaptcha" in lowered or "验证码中间页" in page:
        raise MonitorError("抖音要求验证，暂时无法检查")

    dictionaries: list[dict] = []
    for data in _json_candidates(page):
        dictionaries.extend(_walk(data))
    users = []
    items = []
    for obj in dictionaries:
        candidate = obj.get("user_info") or obj.get("user")
        if isinstance(candidate, dict):
            users.append(candidate)
        if any(key in obj for key in ("unique_id", "sec_uid", "secUid")):
            users.append(obj)
        aweme_id = str(obj.get("aweme_id") or obj.get("awemeId") or "").strip()
        if aweme_id and ("create_time" in obj or "createTime" in obj):
            items.append(obj)

    user = next((item for item in users if _account_matches(item, account_id)), None)
    if user is None and users:
        # 主页以 sec_uid 打开时，页面里的 unique_id 不同；只有唯一账号对象时可采信。
        unique_users = {str(u.get("sec_uid") or u.get("secUid") or u.get("uid") or id(u)): u for u in users}
        if len(unique_users) == 1:
            user = next(iter(unique_users.values()))
    if user is None:
        raise MonitorError("页面未返回可识别的账号数据")
    if user.get("user_canceled") or user.get("is_block") or user.get("status") in {1, 2, 4}:
        raise AccountAbnormalError("账号状态异常或已注销")

    authored = []
    user_sec_uid = str(user.get("sec_uid") or user.get("secUid") or "")
    for item in items:
        author = item.get("author") or {}
        item_sec_uid = str(author.get("sec_uid") or author.get("secUid") or "")
        if not user_sec_uid or not item_sec_uid or item_sec_uid == user_sec_uid:
            authored.append(item)
    latest = max(
        authored,
        key=lambda item: (int(item.get("create_time") or item.get("createTime") or 0), str(item.get("aweme_id") or item.get("awemeId") or "")),
        default={},
    )
    return AccountSnapshot(
        account_id=account_id,
        nickname=str(user.get("nickname") or account_id).strip(),
        latest_aweme_id=str(latest.get("aweme_id") or latest.get("awemeId") or "").strip(),
        latest_create_time=int(latest.get("create_time") or latest.get("createTime") or 0),
        latest_desc=re.sub(r"\s+", " ", str(latest.get("desc") or "")).strip(),
    )


def find_account_sec_uid(page: str, account_id: str) -> str:
    """从搜索结果找出抖音号对应的主页 sec_uid。"""
    for data in _json_candidates(page):
        for obj in _walk(data):
            candidates = [obj]
            for key in ("user", "user_info"):
                if isinstance(obj.get(key), dict):
                    candidates.append(obj[key])
            for user in candidates:
                if _account_matches(user, account_id):
                    return str(user.get("sec_uid") or user.get("secUid") or "").strip()
    return ""


class BrowserAccountClient:
    """复用应用专用 Edge/Chrome，在 HTTP 被验证页拦截时继续监控。"""

    def __init__(self, profile_root):
        from playwright.sync_api import sync_playwright

        self._playwright = sync_playwright().start()
        self._context = None
        errors = []
        for channel in ("msedge", "chrome"):
            try:
                self._context = self._playwright.chromium.launch_persistent_context(
                    user_data_dir=str(profile_root / f"monitor-{channel}"),
                    channel=channel,
                    headless=False,
                    viewport=None,
                    locale="zh-CN",
                    accept_downloads=False,
                )
                break
            except Exception as exc:
                errors.append(f"{channel}: {exc}")
        if self._context is None:
            self._playwright.stop()
            raise MonitorError("无法启动账号监控专用浏览器：" + "；".join(errors))
        self._page = self._context.pages[0] if self._context.pages else self._context.new_page()

    def close(self):
        try:
            self._context.close()
        finally:
            self._playwright.stop()

    def _load(self, url: str) -> str:
        captured = []

        def capture(response):
            content_type = (response.headers.get("content-type") or "").lower()
            if "json" not in content_type:
                return
            try:
                captured.append(response.json())
            except Exception:
                pass

        self._page.on("response", capture)
        try:
            self._page.goto(url, wait_until="domcontentloaded", timeout=30000)
            self._page.wait_for_timeout(4000)
            content = self._page.content()
            try:
                router = self._page.evaluate("() => window._ROUTER_DATA || null")
                if router:
                    captured.append(router)
            except Exception:
                pass
            extra = "".join(
                f'<script type="application/json">{html.escape(json.dumps(data, ensure_ascii=False))}</script>'
                for data in captured
            )
            return content + extra
        finally:
            self._page.remove_listener("response", capture)

    def check(self, account_input: str) -> AccountSnapshot:
        account_id = normalize_account_input(account_input)
        if account_id.startswith("MS4"):
            return parse_account_snapshot(
                self._load(f"https://www.douyin.com/user/{quote(account_id, safe='')}"),
                account_id,
            )
        search_page = self._load(
            f"https://www.douyin.com/search/{quote(account_id, safe='')}?type=user"
        )
        sec_uid = find_account_sec_uid(search_page, account_id)
        if not sec_uid:
            return parse_account_snapshot(search_page, account_id)
        return parse_account_snapshot(
            self._load(f"https://www.douyin.com/user/{quote(sec_uid, safe='')}"),
            sec_uid,
        )


def check_account(account_input: str, session: requests.Session | None = None) -> AccountSnapshot:
    """检查公开页面；搜索抖音号失败时再尝试把输入当作 sec_uid。"""
    account_id = normalize_account_input(account_input)
    own_session = session is None
    client = session or requests.Session()
    client.headers.update({"User-Agent": USER_AGENT, "Referer": "https://www.douyin.com/"})
    urls = [
        f"https://www.douyin.com/search/{quote(account_id, safe='')}?type=user",
        f"https://www.douyin.com/user/{quote(account_id, safe='')}",
    ]
    last_error: Exception | None = None
    try:
        for url in urls:
            try:
                response = client.get(url, timeout=(8, 20))
                if response.status_code in {404, 410}:
                    last_error = AccountAbnormalError("账号主页不存在")
                    continue
                if response.status_code in {403, 429}:
                    raise MonitorError(f"抖音访问受限（HTTP {response.status_code}）")
                response.raise_for_status()
                return parse_account_snapshot(response.text, account_id)
            except AccountAbnormalError as exc:
                last_error = exc
            except (requests.RequestException, MonitorError) as exc:
                last_error = exc
        if isinstance(last_error, AccountAbnormalError):
            raise last_error
        raise MonitorError(str(last_error or "账号检查失败"))
    finally:
        if own_session:
            client.close()
