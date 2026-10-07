# -*- coding: utf-8 -*-

from __future__ import annotations

from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook


LIGHT_DIR = Path(__file__).resolve().parents[1] / "轻量版"
if str(LIGHT_DIR) not in sys.path:
    sys.path.insert(0, str(LIGHT_DIR))

import light_app  # noqa: E402
import light_exporter  # noqa: E402


class LightweightMetadataTests(unittest.TestCase):
    def test_extract_collects_accepts_douyin_statistics_field(self):
        self.assertEqual(
            light_app.extract_collects({"statistics": {"collect_count": "1234"}}),
            1234,
        )
        self.assertEqual(
            light_app.extract_collects({"statistics": {"favorite_count": 56}}),
            56,
        )
        self.assertEqual(light_app.extract_collects({"statistics": {}}), 0)

    def test_extract_author_id_prefers_sec_uid_and_falls_back_to_uid(self):
        self.assertEqual(
            light_app.extract_author_id(
                {"author": {"sec_uid": "MS4wLjABAAA-sec", "uid": "123"}}
            ),
            "MS4wLjABAAA-sec",
        )
        self.assertEqual(
            light_app.extract_author_id({"author": {"uid": "123"}}),
            "123",
        )
        self.assertEqual(light_app.extract_author_id({"author": {}}), "")

    def test_duplicate_filter_keeps_new_url_and_removes_only_duplicate(self):
        duplicate = "https://www.douyin.com/video/123"
        new_url = "https://www.douyin.com/video/456"
        candidate = f"旧文案 {duplicate}；新链接 {new_url}"
        result = light_app.LightDouyinApp._remove_duplicate_urls_from_paste(
            candidate,
            [(duplicate, 1, duplicate)],
        )
        self.assertEqual(result, f"旧文案 ；新链接 {new_url}")

    def test_excel_save_adds_collects_and_preserves_other_content(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "提取记录.xlsx"
            light_exporter.update_records(
                path,
                {
                    1: {
                        "raw_input": "https://www.douyin.com/video/123",
                        "title": "标题",
                        "tags": "#标签",
                        "likes": 10,
                        "collects": 20,
                        "comments": 30,
                        "status": "正常",
                        "aweme_id": "123",
                        "author_id": "MS4wLjABAAA-sec",
                        "updated_at": "now",
                    }
                },
                [1],
            )

            workbook = load_workbook(path)
            workbook.create_sheet("用户工作表")["A1"] = "保留内容"
            sheet = workbook["提取记录"]
            manual_column = sheet.max_column + 1
            sheet.cell(row=1, column=manual_column, value="人工列")
            sheet.cell(row=2, column=manual_column, value="不删除")
            workbook.save(path)
            workbook.close()

            light_exporter.update_records(
                path,
                {
                    1: {
                        "raw_input": "https://www.douyin.com/video/123",
                        "title": "更新标题",
                        "tags": "#更新",
                        "likes": 11,
                        "collects": 22,
                        "comments": 33,
                        "status": "正常",
                        "aweme_id": "123",
                        "author_id": "MS4wLjABAAA-new",
                        "updated_at": "later",
                    }
                },
                [1],
            )

            rows = light_exporter.read_records(path)
            self.assertEqual(rows[1]["collects"], 22)
            self.assertEqual(rows[1]["comments"], 33)
            self.assertEqual(rows[1]["author_id"], "MS4wLjABAAA-new")

            workbook = load_workbook(path, data_only=True)
            self.assertEqual(workbook["用户工作表"]["A1"].value, "保留内容")
            self.assertEqual(workbook["提取记录"].cell(2, manual_column).value, "不删除")
            headers = [cell.value for cell in workbook["提取记录"][1]]
            self.assertIn("收藏数", headers)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
