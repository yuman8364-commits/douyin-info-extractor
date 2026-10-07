# -*- coding: utf-8 -*-
"""轻量版的 Excel 保存层。

只管理轻量版需要的元数据列，但不会重建工作簿：已有工作表、其它列、
已有行和用户自定义内容都会保留。保存时仍采用同目录临时文件加原子替换，
因此不会留下半写入的 ``提取记录.xlsx``。
"""

from __future__ import annotations

import os
from pathlib import Path
import uuid

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


HEADERS = [
    "视频链接",
    "标题",
    "标签",
    "点赞数",
    "收藏数",
    "评论数",
    "状态",
    "顺序",
    "作品ID",
    "作者ID",
    "最后更新",
]
SHEET_NAME = "提取记录"
COLUMN_WIDTHS = {
    "视频链接": 58,
    "标题": 46,
    "标签": 32,
    "点赞数": 12,
    "收藏数": 12,
    "评论数": 12,
    "状态": 28,
    "顺序": 8,
    "作品ID": 24,
    "作者ID": 34,
    "最后更新": 20,
}


class WorkbookInUseError(PermissionError):
    """目标工作簿被 Excel/WPS 占用。"""


def _headers(sheet) -> list[str]:
    if sheet.max_row < 1:
        return []
    return [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in sheet[1]
    ]


def _looks_managed(sheet) -> bool:
    headers = set(_headers(sheet))
    return "视频链接" in headers and "顺序" in headers


def _select_sheet(workbook, *, create: bool):
    """优先选择本工具工作表；未知工作表不被覆盖。"""
    if SHEET_NAME in workbook.sheetnames:
        sheet = workbook[SHEET_NAME]
        if _looks_managed(sheet) or sheet.max_row <= 1:
            return sheet

    for sheet in workbook.worksheets:
        if _looks_managed(sheet):
            return sheet

    if not create:
        return None

    if len(workbook.worksheets) == 1:
        sheet = workbook.active
        if sheet.max_row <= 1 and not any(_headers(sheet)):
            sheet.title = SHEET_NAME
            return sheet

    title = SHEET_NAME
    suffix = 2
    while title in workbook.sheetnames:
        title = f"{SHEET_NAME}{suffix}"
        suffix += 1
    return workbook.create_sheet(title)


def _ensure_headers(sheet) -> dict[str, int]:
    current = _headers(sheet)
    if not any(current):
        current = []

    # 只追加缺少的管理列，不删除或重排已有列。
    for name in HEADERS:
        if name not in current:
            current.append(name)

    for column, text in enumerate(current, 1):
        cell = sheet.cell(row=1, column=column, value=text)
        if text in HEADERS:
            cell.font = Font(bold=True)
            sheet.column_dimensions[get_column_letter(column)].width = (
                COLUMN_WIDTHS.get(text, 16)
            )
    return {name: index + 1 for index, name in enumerate(current) if name}


def _parse_seq(value) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        try:
            number = float(value)
        except (TypeError, ValueError):
            return None
        return int(number) if number.is_integer() and number > 0 else None
    text = str(value).strip()
    if text.isdigit() and int(text) > 0:
        return int(text)
    return None


def _as_int(value) -> int:
    if value is None or isinstance(value, bool):
        return 0
    if isinstance(value, (int, float)):
        try:
            return max(0, int(value))
        except (TypeError, ValueError, OverflowError):
            return 0
    text = str(value).strip().replace(",", "")
    try:
        return max(0, int(float(text)))
    except (TypeError, ValueError, OverflowError):
        return 0


def _record_from_row(values: list, layout: dict[str, int]) -> tuple[int | None, dict]:
    def pick(name: str):
        column = layout.get(name)
        index = column - 1 if column else None
        return values[index] if index is not None and index < len(values) else None

    seq = _parse_seq(pick("顺序"))
    record = {
        "raw_input": str(pick("视频链接") or ""),
        "title": str(pick("标题") or ""),
        "tags": str(pick("标签") or ""),
        "likes": _as_int(pick("点赞数")),
        "collects": _as_int(pick("收藏数")),
        "comments": _as_int(pick("评论数")),
        "status": str(pick("状态") or ""),
        "aweme_id": str(pick("作品ID") or "").strip(),
        "author_id": str(pick("作者ID") or "").strip(),
        "updated_at": str(pick("最后更新") or ""),
    }
    return seq, record


def read_records(path) -> dict[int, dict]:
    """读取轻量版管理的列；未知工作簿结构只读并返回空字典。"""
    path = Path(path)
    if not path.exists():
        return {}

    workbook = load_workbook(path, read_only=False)
    try:
        sheet = _select_sheet(workbook, create=False)
        if sheet is None:
            return {}
        layout = {
            name: index + 1
            for index, name in enumerate(_headers(sheet))
            if name
        }
        result: dict[int, dict] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            seq, record = _record_from_row(list(row), layout)
            if seq is not None:
                result[seq] = record
        return result
    finally:
        workbook.close()


def _write_record(sheet, row: int, record: dict, seq: int, layout: dict[str, int]) -> None:
    values = {
        "视频链接": record.get("raw_input") or "",
        "标题": record.get("title") or "",
        "标签": record.get("tags") or "",
        "点赞数": _as_int(record.get("likes")),
        "收藏数": _as_int(record.get("collects", record.get("favorites"))),
        "评论数": _as_int(record.get("comments")),
        "状态": record.get("status") or "",
        "顺序": seq,
        "作品ID": str(record.get("aweme_id") or ""),
        "作者ID": str(record.get("author_id") or ""),
        "最后更新": record.get("updated_at") or "",
    }
    for name, value in values.items():
        cell = sheet.cell(row=row, column=layout[name], value=value)
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=name in {"视频链接", "标题", "标签", "状态"},
        )


def _save_atomic(workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp{path.suffix}")
    try:
        workbook.save(temp)
        check = load_workbook(temp, read_only=True)
        check.close()
        os.replace(temp, path)
    except PermissionError as exc:
        raise WorkbookInUseError(
            "提取记录.xlsx 正被 Excel/WPS 占用，请关闭该工作簿后重试"
        ) from exc
    finally:
        temp.unlink(missing_ok=True)


def update_records(
    path,
    records: dict[int, dict],
    seqs: list[int] | None = None,
) -> None:
    """增量写入元数据，并保留工作簿中的其它内容。"""
    path = Path(path)
    workbook = load_workbook(path) if path.exists() else Workbook()
    try:
        sheet = _select_sheet(workbook, create=True)
        layout = _ensure_headers(sheet)
        seq_column = layout["顺序"]
        row_by_seq: dict[int, int] = {}
        for row in range(2, sheet.max_row + 1):
            seq = _parse_seq(sheet.cell(row=row, column=seq_column).value)
            if seq is not None:
                row_by_seq[seq] = row

        selected = seqs if seqs is not None else sorted(records)
        for seq in selected:
            record = records.get(seq)
            if record is None:
                continue
            row = row_by_seq.get(seq)
            if row is None:
                row = sheet.max_row + 1
                row_by_seq[seq] = row
            _write_record(sheet, row, record, int(seq), layout)
        _save_atomic(workbook, path)
    finally:
        workbook.close()


def append_record(path, record: dict, seq: int) -> int:
    """测试和小型调用方使用的单条新增兼容接口。"""
    update_records(path, {int(seq): record}, [int(seq)])
    rows = read_records(path)
    return sorted(rows).index(int(seq)) + 2 if int(seq) in rows else 0
